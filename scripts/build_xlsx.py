import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV_PATH = os.path.join(ROOT, "colony_numeric_dataset.csv")
OUT_PATH = os.path.join(
    os.path.dirname(ROOT), "colony_classification_project", "data", "colony_numeric_dataset.xlsx"
)

FEATURE_COLS = [
    "equiv_diameter", "circularity", "solidity",
    "mean_hue", "mean_sat", "mean_val",
    "std_hue", "std_sat", "std_val",
    "texture_std",
]
STRAINS = ["HS2", "ALE", "HMF", "HS19"]

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4C72B0")
BODY_FONT = Font(name="Arial")
BAND_FILL = PatternFill("solid", fgColor="F2F2F2")


def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autofit(ws, ncols, min_width=10, max_width=45):
    for c in range(1, ncols + 1):
        col_letter = get_column_letter(c)
        max_len = max(
            (len(str(ws.cell(row=r, column=c).value)) for r in range(1, ws.max_row + 1)
             if ws.cell(row=r, column=c).value is not None),
            default=min_width,
        )
        ws.column_dimensions[col_letter].width = min(max(max_len + 2, min_width), max_width)


def band_rows(ws, first_data_row, last_data_row, ncols):
    for r in range(first_data_row, last_data_row + 1):
        if (r - first_data_row) % 2 == 1:
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).fill = BAND_FILL


def main():
    rows = list(csv.DictReader(open(CSV_PATH, encoding="utf-8")))
    plate_groups = sorted(set(r["plate_group"] for r in rows))

    wb = Workbook()

    # ---- Sheet 1: Colony Data ----
    ws = wb.active
    ws.title = "Colony Data"
    headers = ["strain", "plate_group", "crop_path", "source_image"] + FEATURE_COLS
    ws.append(headers)
    for r in rows:
        ws.append([r[h] if h in ("strain", "plate_group", "crop_path", "source_image") else float(r[h])
                   for h in headers])
    n_data = len(rows)
    for r in range(2, n_data + 2):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = BODY_FONT
    style_header(ws, len(headers))
    ws.freeze_panes = "A2"
    band_rows(ws, 2, n_data + 1, len(headers))
    autofit(ws, len(headers))

    data_first_row, data_last_row = 2, n_data + 1
    strain_col = "A"
    plate_col = "B"
    feat_col_letter = {name: get_column_letter(5 + i) for i, name in enumerate(FEATURE_COLS)}
    header_row_range = f"'Colony Data'!$E$1:$N$1"

    # ---- Sheet 2: Strain Summary ----
    ws2 = wb.create_sheet("Strain Summary")
    ws2.append(["strain", "feature", "n", "sum_x", "sum_x2", "mean", "std"])
    style_header(ws2, 7)
    row = 2
    for strain in STRAINS:
        for feat in FEATURE_COLS:
            fcol = feat_col_letter[feat]
            data_rng = f"'Colony Data'!${fcol}${data_first_row}:${fcol}${data_last_row}"
            strain_rng = f"'Colony Data'!${strain_col}${data_first_row}:${strain_col}${data_last_row}"
            ws2.cell(row=row, column=1, value=strain)
            ws2.cell(row=row, column=2, value=feat)
            ws2.cell(row=row, column=3, value=f"=SUMPRODUCT(({strain_rng}=$A{row})*1)")
            ws2.cell(row=row, column=4, value=f"=SUMPRODUCT(({strain_rng}=$A{row})*({data_rng}))")
            ws2.cell(row=row, column=5, value=f"=SUMPRODUCT(({strain_rng}=$A{row})*({data_rng})^2)")
            ws2.cell(row=row, column=6, value=f"=IF(C{row}=0,\"\",D{row}/C{row})")
            ws2.cell(row=row, column=7,
                     value=f"=IF(C{row}<=1,\"\",SQRT((E{row}-(D{row}^2)/C{row})/(C{row}-1)))")
            for c in range(1, 8):
                ws2.cell(row=row, column=c).font = BODY_FONT
            row += 1
    n2 = row - 1
    ws2.freeze_panes = "A2"
    band_rows(ws2, 2, n2, 7)
    autofit(ws2, 7)

    # ---- Sheet 3: Plate Summary ----
    ws3 = wb.create_sheet("Plate Summary")
    ws3.append(["plate_group", "strain", "feature", "n", "sum_x", "sum_x2", "mean", "std"])
    style_header(ws3, 8)
    plate_to_strain = {r["plate_group"]: r["strain"] for r in rows}
    plate_groups_sorted = sorted(plate_groups, key=lambda g: (plate_to_strain[g], g))
    row = 2
    for grp in plate_groups_sorted:
        strain = plate_to_strain[grp]
        for feat in FEATURE_COLS:
            fcol = feat_col_letter[feat]
            data_rng = f"'Colony Data'!${fcol}${data_first_row}:${fcol}${data_last_row}"
            plate_rng = f"'Colony Data'!${plate_col}${data_first_row}:${plate_col}${data_last_row}"
            ws3.cell(row=row, column=1, value=grp)
            ws3.cell(row=row, column=2, value=strain)
            ws3.cell(row=row, column=3, value=feat)
            ws3.cell(row=row, column=4, value=f"=SUMPRODUCT(({plate_rng}=$A{row})*1)")
            ws3.cell(row=row, column=5, value=f"=SUMPRODUCT(({plate_rng}=$A{row})*({data_rng}))")
            ws3.cell(row=row, column=6, value=f"=SUMPRODUCT(({plate_rng}=$A{row})*({data_rng})^2)")
            ws3.cell(row=row, column=7, value=f"=IF(D{row}=0,\"\",E{row}/D{row})")
            ws3.cell(row=row, column=8,
                     value=f"=IF(D{row}<=1,\"\",SQRT((F{row}-(E{row}^2)/D{row})/(D{row}-1)))")
            for c in range(1, 9):
                ws3.cell(row=row, column=c).font = BODY_FONT
            row += 1
    n3 = row - 1
    ws3.freeze_panes = "A2"
    band_rows(ws3, 2, n3, 8)
    autofit(ws3, 8)

    os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
    wb.save(OUT_PATH)
    print("wrote", OUT_PATH)


if __name__ == "__main__":
    main()
